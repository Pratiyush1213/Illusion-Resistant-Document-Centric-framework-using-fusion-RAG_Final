import streamlit as st
import re
import io

from langchain_core.documents import Document as LCDocument
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_community.vectorstores import FAISS
from langchain_community.embeddings.fastembed import FastEmbedEmbeddings
from docx import Document
from pypdf import PdfReader

# ── Optional OCR ──
try:
    import pytesseract
    from PIL import Image
    OCR_AVAILABLE = True
except ImportError:
    OCR_AVAILABLE = False

# ── Page config ──
st.set_page_config(page_title="FusionRAG · Doc Chat", layout="wide", page_icon="⚡")

# ── CSS ──
st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=Syne:wght@400;700;800&family=Inter:wght@300;400;500&display=swap');

html, body, [class*="css"] { font-family: 'Inter', sans-serif; }

.stApp {
    background: #07060f;
    color: #e8e6f0;
}
.stApp::before {
    content: '';
    position: fixed;
    inset: 0;
    background:
        radial-gradient(ellipse 80% 60% at 20% 10%, rgba(120,40,200,0.18) 0%, transparent 60%),
        radial-gradient(ellipse 60% 50% at 80% 80%, rgba(60,20,140,0.14) 0%, transparent 60%),
        radial-gradient(ellipse 40% 40% at 60% 20%, rgba(200,80,255,0.07) 0%, transparent 50%);
    pointer-events: none;
    z-index: 0;
}
h1 {
    font-family: 'Syne', sans-serif !important;
    font-size: 3rem !important;
    font-weight: 800 !important;
    background: linear-gradient(135deg, #ffffff 30%, #b57bee 100%);
    -webkit-background-clip: text;
    -webkit-text-fill-color: transparent;
    background-clip: text;
    text-align: center;
    letter-spacing: -1px;
    margin-bottom: 0.2rem !important;
    padding-top: 1.5rem;
}
.subtitle {
    text-align: center;
    color: #7c6fa0;
    font-size: 0.95rem;
    margin-bottom: 2rem;
    letter-spacing: 0.04em;
}
section[data-testid="stSidebar"] {
    background: #0e0b1a !important;
    border-right: 1px solid #1e1530;
}
section[data-testid="stSidebar"] * { color: #c8bfe0 !important; }
section[data-testid="stSidebar"] h1,
section[data-testid="stSidebar"] h2,
section[data-testid="stSidebar"] h3 {
    color: #ffffff !important;
    font-family: 'Syne', sans-serif !important;
}
[data-testid="stFileUploader"] {
    border: 1.5px dashed #3d2b6b !important;
    border-radius: 14px !important;
    background: rgba(61,43,107,0.12) !important;
}
.stTextInput input {
    background: #110e1f !important;
    color: #e8e6f0 !important;
    border: 1.5px solid #2d2050 !important;
    border-radius: 12px !important;
    padding: 0.75rem 1rem !important;
    font-size: 1rem !important;
    transition: border-color 0.3s, box-shadow 0.3s;
}
.stTextInput input:focus {
    border-color: #7c3aed !important;
    box-shadow: 0 0 0 3px rgba(124,58,237,0.18) !important;
}
.stTextInput label { color: #9d86c8 !important; font-size: 0.85rem !important; }
.stSlider label { color: #9d86c8 !important; }

.answer-card {
    background: linear-gradient(145deg, #130f24, #1a1330);
    border: 1px solid #2a1f4a;
    border-radius: 18px;
    padding: 1.5rem 1.75rem;
    margin: 1rem 0;
    position: relative;
    overflow: hidden;
    transition: transform 0.25s ease, box-shadow 0.25s ease;
    animation: cardIn 0.5s ease backwards;
}
.answer-card::before {
    content: '';
    position: absolute;
    top: 0; left: 0; right: 0;
    height: 2px;
    background: linear-gradient(90deg, #7c3aed, #a855f7, #7c3aed);
    background-size: 200% 100%;
    animation: shimmer 3s linear infinite;
}
.answer-card:hover {
    transform: translateY(-3px);
    box-shadow: 0 12px 40px rgba(124,58,237,0.2);
}
@keyframes cardIn {
    from { opacity: 0; transform: translateY(16px); }
    to   { opacity: 1; transform: translateY(0); }
}
@keyframes shimmer {
    0%   { background-position: 200% 0; }
    100% { background-position: -200% 0; }
}
.card-label {
    font-family: 'Syne', sans-serif;
    font-size: 0.7rem;
    font-weight: 700;
    letter-spacing: 0.12em;
    text-transform: uppercase;
    color: #7c3aed;
    margin-bottom: 0.6rem;
}
.card-body {
    font-size: 0.97rem;
    line-height: 1.85;
    color: #ccc4e0;
}
.hl {
    background: linear-gradient(90deg, rgba(124,58,237,0.35), rgba(168,85,247,0.25));
    border-radius: 4px;
    padding: 1px 4px;
    color: #e0d4ff;
    font-weight: 500;
}
.empty-state {
    text-align: center;
    padding: 4rem 2rem;
    color: #3d3060;
}
.empty-state .icon { font-size: 3.5rem; margin-bottom: 1rem; }
hr { border-color: #1e1530 !important; margin: 1.5rem 0 !important; }
</style>
""", unsafe_allow_html=True)

# ── Header ──
st.markdown("<h1>⚡ FusionRAG</h1>", unsafe_allow_html=True)
st.markdown('<p class="subtitle">Upload documents · Ask questions · Get instant answers</p>', unsafe_allow_html=True)

# ── Sidebar ──
with st.sidebar:
    st.markdown("### 📂 Documents")
    accepted = ["pdf", "docx"]
    if OCR_AVAILABLE:
        accepted += ["png", "jpg", "jpeg"]

    files = st.file_uploader(
        "Upload PDFs or DOCX files",
        type=accepted,
        accept_multiple_files=True
    )
    st.divider()
    st.markdown("### ⚙️ Settings")
    top_k = st.slider("Results to show", 1, 10, 3)

# ── Extractors ──
def extract_pdf(file) -> list[LCDocument]:
    """Read PDF directly in memory — no temp file needed."""
    reader = PdfReader(io.BytesIO(file.read()))
    docs = []
    for i, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        if text.strip():
            docs.append(LCDocument(
                page_content=text,
                metadata={"source": file.name, "page": i + 1}
            ))
    return docs

def extract_docx(file) -> list[LCDocument]:
    doc = Document(file)
    text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    if text.strip():
        return [LCDocument(page_content=text, metadata={"source": file.name})]
    return []

def extract_image(file) -> list[LCDocument]:
    img = Image.open(file)
    text = pytesseract.image_to_string(img)
    if text.strip():
        return [LCDocument(page_content=text, metadata={"source": file.name})]
    return []

def highlight_query(text, query):
    for word in re.findall(r'\w+', query):
        text = re.sub(rf'(?i)({re.escape(word)})', r'<span class="hl">\1</span>', text)
    return text

# ── Session state ──
for key in ["db", "last_file_names"]:
    if key not in st.session_state:
        st.session_state[key] = None if key == "db" else []

# ── Detect file changes → reset index ──
current_names = sorted([f.name for f in files]) if files else []
if current_names != st.session_state.last_file_names:
    st.session_state.db = None
    st.session_state.last_file_names = current_names

# ── Build index ──
if files and st.session_state.db is None:
    progress = st.progress(0, text="Reading files…")
    raw_docs = []

    for i, file in enumerate(files):
        ext = file.name.rsplit(".", 1)[-1].lower()
        progress.progress(int((i / len(files)) * 40), text=f"Reading {file.name}…")
        try:
            if ext == "pdf":
                raw_docs.extend(extract_pdf(file))
            elif ext == "docx":
                raw_docs.extend(extract_docx(file))
            elif ext in ["png", "jpg", "jpeg"] and OCR_AVAILABLE:
                raw_docs.extend(extract_image(file))
        except Exception as e:
            st.warning(f"Could not process **{file.name}**: {e}")

    if not raw_docs:
        progress.empty()
        st.error("No readable text found in the uploaded files.")
        st.stop()

    progress.progress(50, text="Splitting into chunks…")
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=800,
        chunk_overlap=150,
        separators=["\n\n", "\n", ". ", "? ", "! ", " "]
    )
    chunks = splitter.split_documents(raw_docs)

    progress.progress(70, text="Loading embedding model…")
    embeddings = FastEmbedEmbeddings(model_name="BAAI/bge-small-en-v1.5")

    progress.progress(85, text="Building search index…")
    st.session_state.db = FAISS.from_documents(chunks, embeddings)

    progress.progress(100, text="Done!")
    progress.empty()
    st.success(f"✅ Ready! Indexed {len(chunks)} chunks from {len(files)} file(s).")

# ── Main area ──
if not files:
    st.markdown("""
    <div class="empty-state">
        <div class="icon">📄</div>
        <p>Upload PDF or DOCX files from the sidebar to get started.</p>
    </div>
    """, unsafe_allow_html=True)
else:
    query = st.text_input(
        "Ask a question from your documents",
        placeholder="e.g. What is the main topic of the document?"
    )

    if query and st.session_state.db:
        results = st.session_state.db.similarity_search(query, k=top_k)
        st.markdown("<hr>", unsafe_allow_html=True)

        if not results:
            st.warning("No relevant content found. Try rephrasing your question.")
        else:
            for i, res in enumerate(results, 1):
                sentences = re.split(r'(?<=[.!?])\s+', res.page_content.strip())
                answer = " ".join(sentences[:6])
                answer = highlight_query(answer, query)

                src = res.metadata.get("source", "")
                page = res.metadata.get("page", "")
                src_info = f"{src} · p.{page}" if page else src
                src_tag = f'<span style="float:right;color:#4b3880;font-size:0.75rem;">{src_info}</span>' if src_info else ""

                st.markdown(f"""
                <div class="answer-card" style="animation-delay:{(i-1)*0.08}s">
                    <div class="card-label">Result {i} {src_tag}</div>
                    <div class="card-body">{answer}</div>
                </div>
                """, unsafe_allow_html=True)
